from pathlib import Path
import sys,json
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
import reportes
from pypdf import PdfReader
from openpyxl import load_workbook
out=Path(__file__).resolve().parent
res=[]
for tipo in reportes.TIPOS:
 r=reportes.obtener_datos_reporte(tipo)
 pdf=reportes.exportar_pdf(r,out/reportes.nombre_archivo(r,'pdf'))
 xlsx=reportes.exportar_excel(r,out/reportes.nombre_archivo(r,'xlsx'))
 wb=load_workbook(xlsx)
 res.append({'tipo':tipo,'filas':len(r['filas']),'pdf':pdf.name,'xlsx':xlsx.name,'paginas':len(PdfReader(pdf).pages),'hojas':wb.sheetnames})
 print(tipo,len(r['filas']),res[-1]['paginas'])
(out/'verificacion.json').write_text(json.dumps(res,ensure_ascii=False,indent=2),encoding='utf-8')
