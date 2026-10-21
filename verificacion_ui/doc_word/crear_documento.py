from pathlib import Path
import sys
import json, textwrap
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
ROOT=BASE.parents[1]
E=json.loads((BASE/'evidencia_sql.json').read_text(encoding='utf-8'))
D=Document()
sec=D.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.54)
sec.different_first_page_header_footer=True
sec.footer_distance=Cm(1.25)
for name,size,bold in [('Normal',11,False),('Title',16,True),('Heading 1',16,True),('Heading 2',13,True),('Heading 3',12,True),('Caption',11,False)]:
    st=D.styles[name]; st.font.name='Arial'; st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'),'Arial')
    fmt=st.paragraph_format; fmt.line_spacing=1.5; fmt.space_after=Pt(6)
    if name.startswith('Heading'):
        fmt.space_before=Pt(12); fmt.space_after=Pt(7); fmt.keep_with_next=True
    if name=='Normal': fmt.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY; fmt.widow_control=True
    if name=='Caption': fmt.alignment=WD_ALIGN_PARAGRAPH.CENTER; fmt.line_spacing=1.15; fmt.space_after=Pt(10)
    if name=='Title':
        bd=OxmlElement('w:pBdr')
        for edge in ('top','left','bottom','right','between','bar'):
            e=OxmlElement('w:'+edge); e.set(qn('w:val'),'nil'); bd.append(e)
        st.element.get_or_add_pPr().append(bd)
for n in ('TOC 1','TOC 2'):
    if n not in D.styles: D.styles.add_style(n,1)
    st=D.styles[n]; st.font.name='Arial'; st.font.size=Pt(11); st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.line_spacing=1.15; st.paragraph_format.space_after=Pt(3)

# Multilevel, editable Word numbering tied to native Heading styles.
nums=D.part.numbering_part.element
ab=OxmlElement('w:abstractNum'); ab.set(qn('w:abstractNumId'),'70')
typ=OxmlElement('w:multiLevelType'); typ.set(qn('w:val'),'multilevel'); ab.append(typ)
for level in range(3):
    lv=OxmlElement('w:lvl'); lv.set(qn('w:ilvl'),str(level))
    for tag,val in [('start','1'),('numFmt','decimal'),('pStyle',f'Heading{level+1}'),('lvlText','.'.join('%'+str(i+1) for i in range(level+1))+'.'),('suff','space')]:
        x=OxmlElement('w:'+tag); x.set(qn('w:val'),val); lv.append(x)
    ab.append(lv)
nums.append(ab)
nu=OxmlElement('w:num'); nu.set(qn('w:numId'),'70'); a=OxmlElement('w:abstractNumId'); a.set(qn('w:val'),'70'); nu.append(a); nums.append(nu)
for level in range(3):
    ppr=D.styles[f'Heading {level+1}'].element.get_or_add_pPr(); np=OxmlElement('w:numPr')
    il=OxmlElement('w:ilvl'); il.set(qn('w:val'),str(level)); ni=OxmlElement('w:numId'); ni.set(qn('w:val'),'70'); np.extend([il,ni]); ppr.append(np)

def field(p,code):
    for typ in ['begin','instruction','end']:
        r=OxmlElement('w:r')
        if typ=='instruction':
            el=OxmlElement('w:instrText'); el.set(qn('xml:space'),'preserve'); el.text=code
        else: el=OxmlElement('w:fldChar'); el.set(qn('w:fldCharType'),typ)
        r.append(el); p._p.append(r)
foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; field(foot,' PAGE ')
settings=D.settings.element
up=OxmlElement('w:updateFields'); up.set(qn('w:val'),'true'); settings.append(up)
lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'es-NI'); D.styles['Normal'].element.get_or_add_rPr().append(lang)
D.core_properties.title='Sistema Inteligente para la Gestión y Análisis de Inventario en Almacenes mediante Inteligencia Artificial'
D.core_properties.author='Alexander Josué Lorio Paizano; Andrea Sofia Ramírez Centeno; Roger Martin Pérez Castillo; Guadalupe del Carmen Tellez Palacio'
D.core_properties.subject='Informe académico actualizado: metodología CRISP-DM y estado funcional'
D.core_properties.keywords='Inventario; CRISP-DM; SQL Server; JARVIS; YuNet; SFace'

def p(s,bold=False):
    x=D.add_paragraph(); x.add_run(s).bold=bold; return x
def h(s,level=1): return D.add_heading(s,level)
def bullet(s):
    x=D.add_paragraph(s,style='List Bullet'); x.paragraph_format.line_spacing=1.5; return x
T=0; F=0
def table(title,heads,rows,widths=None):
    global T
    T+=1; cap=D.add_paragraph(f'Tabla {T}. {title}',style='Caption'); cap.paragraph_format.keep_with_next=True
    t=D.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    if widths is None: widths=[15.92/len(heads)]*len(heads)
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    for cell,s,w in zip(t.rows[0].cells,heads,widths): cell.text=str(s); cell.width=Cm(w)
    trpr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
    for row in rows:
        cs=t.add_row().cells
        for c,s,w in zip(cs,row,widths): c.text=str(s); c.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
    pr.append(borders)
    for idx,row in enumerate(t.rows):
        no=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no)
        for c in row.cells:
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr(); margin=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge); el.set(qn('w:w'),'85'); el.set(qn('w:type'),'dxa'); margin.append(el)
            cp.append(margin)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'20364B' if idx==0 else ('F3F5F7' if idx%2==0 else 'FFFFFF')); cp.append(shade)
            for para in c.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT; para.paragraph_format.line_spacing=1.15; para.paragraph_format.space_after=Pt(3)
                if idx==0: para.paragraph_format.keep_with_next=True
                for run in para.runs:
                    run.font.size=Pt(11); run.font.name='Arial'
                    if idx==0: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    D.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def diagram(name,items):
    font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',32)
    W=1600; H=290; im=Image.new('RGB',(W,H),'white'); dr=ImageDraw.Draw(im)
    n=len(items); bw=int((W-90-(n-1)*35)/n); y=40
    for i,label in enumerate(items):
        x=35+i*(bw+35); dr.rounded_rectangle((x,y,x+bw,y+190),radius=12,fill='#EFF4F7',outline='#20364B',width=3)
        lines=[]
        for line in label.split('\n'):
            current=''
            for word in line.split():
                trial=(current+' '+word).strip()
                if dr.textbbox((0,0),trial,font=font)[2]>bw-22 and current:
                    lines.append(current); current=word
                else: current=trial
            if current: lines.append(current)
        th=len(lines)*39; ty=y+(190-th)/2
        for line in lines:
            box=dr.textbbox((0,0),line,font=font); dr.text((x+(bw-(box[2]-box[0]))/2,ty),line,font=font,fill='#142B3C'); ty+=39
        if i<n-1:
            ax=x+bw+4; ay=y+95; dr.line((ax,ay,ax+27,ay),fill='#168E94',width=5); dr.polygon([(ax+27,ay),(ax+17,ay-9),(ax+17,ay+9)],fill='#168E94')
    path=BASE/(name+'.png'); im.save(path); return path
def figure(path,title,width=15.7):
    global F
    F+=1; x=D.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.keep_with_next=True; x.paragraph_format.line_spacing=1
    run=x.add_run(); run.add_picture(str(path),width=Cm(width))
    pic=run._r.xpath('.//wp:docPr')[0]; pic.set('descr',title)
    D.add_paragraph(f'Figura {F}. {title}',style='Caption')

# Cover maintains the institutional information and team from the original.
x=D.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.add_run().add_picture(str(BASE/'ref1_0.png'),width=Cm(6))
for text,size,bold in [
 ('Área de Conocimiento de Tecnología de la Información y Comunicación',11,True),
 ('Departamento de Ingeniería de Sistemas y Telecomunicaciones',11,False),
 ('Programa de Ingeniería de Sistemas',11,False),
 ('Proyecto de Curso de Inteligencia Artificial',13,True),
 (D.core_properties.title,16,True),
 ('Informe académico actualizado',12,False),
 ('Asignatura: Inteligencia Artificial\nGrupo: 4T4-SIS-S',11,False),
 ('Docente: PhD. Denis Eduardo Hernández García',11,False),
 ('Elaborado por:\nBr. Alexander Josué Lorio Paizano\nBr. Andrea Sofia Ramírez Centeno\nBr. Roger Martin Pérez Castillo\nBr. Guadalupe del Carmen Tellez Palacio',11,False),
 ('3 de octubre de 2026\nManagua, Nicaragua',11,False)]:
    x=D.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.paragraph_format.line_spacing=1.15; x.paragraph_format.space_after=Pt(10)
    run=x.add_run(text); run.font.size=Pt(size); run.bold=bold
D.add_page_break()
x=D.add_paragraph('Índice',style='Title'); x.alignment=WD_ALIGN_PARAGRAPH.LEFT
field(D.add_paragraph(),' TOC \\o "1-2" \\h \\z \\u ')
D.add_page_break()

h('Introducción')
p('El almacén de un taller de mecánica reúne repuestos, herramientas, lubricantes, filtros y materiales necesarios para mantenimiento y reparación. Conocer su disponibilidad exige registrar las entradas y salidas, mantener su ubicación y relacionar cada operación con un empleado responsable. Cuando esos registros están dispersos, las consultas y la revisión de diferencias se vuelven más difíciles. Este problema conserva el enfoque planteado en los entregables iniciales del proyecto [1, 2].')
p('El Sistema Inteligente para la Gestión y Análisis de Inventario en Almacenes mediante Inteligencia Artificial integra esa operación en una aplicación de escritorio. La base de datos concentra productos, existencias por ubicación, movimientos y empleados; sobre esos datos se incorporan consultas en lenguaje natural, reconocimiento facial, estimaciones de cobertura y detección de cantidades inusuales. El propósito es apoyar al responsable del almacén con información verificable, manteniendo su intervención en las decisiones.')
p('La versión revisada al 3 de octubre de 2026 dispone de una interfaz funcional y un escenario de demostración identificado. Los resultados operativos proceden de SQL Server y de las comprobaciones existentes. La demostración permite observar los flujos, pero no constituye una evaluación del consumo real del taller ni de la precisión facial en condiciones de campo. Esta diferencia determina el alcance de los resultados y de las conclusiones.')

h('Descripción general del proyecto')
p('La aplicación permite administrar el catálogo de productos, registrar operaciones, localizar existencias y consultar historial. Su unidad operativa es la combinación producto–ubicación: un artículo puede almacenarse en varias ubicaciones y su stock total corresponde a la suma de sus cantidades vigentes. Las transferencias redistribuyen el inventario y conservan su cantidad global.')
p('JARVIS funciona como punto de consulta común para inventario, ubicaciones, empleados, movimientos, predicción y anomalías. Sus respuestas escritas y habladas provienen de los mismos servicios utilizados por las pantallas. La visión facial identifica empleados inscritos cuando el pipeline y su calibración son válidos; no establece cuentas de acceso ni sustituye la selección del responsable de una operación.')
table('Componentes y situación funcional',['Componente','Situación actual','Validación restante'],[
 ('Inventario y movimientos','IMPLEMENTADA: registro transaccional, stock por ubicación, lotes y conteos.','Uso sostenido con registros del almacén.'),
 ('Mapa lógico','IMPLEMENTADA: consulta SQL y localización de múltiples ubicaciones.','Correspondencia del catálogo con el espacio físico.'),
 ('JARVIS escrito y voz','IMPLEMENTADA: procesamiento compartido y acciones de consulta.','Pruebas acústicas y de equipos en el taller.'),
 ('Reconocimiento facial','IMPLEMENTADA la infraestructura; EN PROCESO la calibración operativa.','PENDIENTE DE VALIDACIÓN: precisión y alcance físico.'),
 ('Predicción y anomalías','IMPLEMENTADA: motores estadísticos explicables.','PENDIENTE DE VALIDACIÓN con historia observada.'),
 ('LEDs / ESP32','No existe integración funcional en el código revisado.','Trabajo futuro; no forma parte de los resultados.')],[3.3,7,5.62])
p('El programa inicia con python main.py y abre el inventario directamente. Los empleados son entidades operativas relacionadas con registros y muestras faciales. Esta versión no contiene administración de usuarios, roles, permisos, login, auditoría ni integración de identidad externa.')

h('Metodología CRISP-DM')
p('CRISP-DM mantiene la secuencia de comprensión del negocio, comprensión de los datos, preparación, modelado, evaluación y despliegue descrita en los antecedentes [1, 2]. En el proyecto estas fases se relacionan de manera iterativa: una insuficiencia de datos o un error de identificación puede requerir revisar la captura, las reglas de calidad y el modelo antes de continuar con su evaluación.')
table('Relación entre CRISP-DM y el avance actual',['Fase','Evidencia actual','Alcance del avance'],[
 ('Comprensión del negocio','Problema, objetivos, alcance e interesados definidos en los entregables.','Base académica conservada y vinculada a la versión actual.'),
 ('Comprensión de los datos','Inspección de tablas, relaciones, cantidades y restricciones SQL.','Caracterización disponible; la historia principal es simulada.'),
 ('Preparación de los datos','Validaciones operativas, exclusión de registros no vigentes, series de salidas y control de calidad facial.','Procesos implementados; ampliar datos observados sigue pendiente.'),
 ('Modelado','YuNet/SFace, NLP estructurado, promedio diario y comparación estadística de cantidades.','Componentes integrados; calibración facial en proceso.'),
 ('Evaluación','Pruebas funcionales, integridad SQL y revisión de pantallas.','No acredita métricas experimentales de IA ni resultados de campo.'),
 ('Despliegue','Ejecución local mediante main.py y SQL Server.','Prototipo de escritorio disponible; operación prolongada pendiente.')],[3.2,7.1,5.62])
p('La evaluación funcional comprueba que los servicios y pantallas cumplen sus reglas. La evaluación analítica debe contrastar estimaciones con consumo posterior y validar identificaciones con muestras independientes. Por tanto, las seis fases tienen actividades desarrolladas, pero no se declara un cierre completo del ciclo experimental ni una implantación definitiva en el almacén.')

h('Fase 1. Comprensión del negocio')
h('Definición del problema',2)
p('Los antecedentes describen dificultades para conocer existencias oportunamente, localizar productos y revisar diferencias entre cantidades registradas y físicas. También requieren identificar el empleado responsable de cada movimiento y aprovechar el historial para anticipar agotamientos o señalar operaciones que necesitan revisión. El problema se concentra en organizar, controlar y analizar información del almacén, no en automatizar todas sus decisiones [2].')
h('Antecedentes y justificación',2)
p('Los registros manuales y las hojas de cálculo permiten conservar información básica, pero su actualización depende de procedimientos constantes y no siempre conectan producto, ubicación, movimiento y responsable. El proyecto reúne esas relaciones en SQL Server y expone consultas mediante la interfaz y JARVIS. Su aporte académico consiste en articular programación, bases de datos, visión artificial, lenguaje natural y análisis histórico en una necesidad concreta [1].')
p('La centralización proporciona una base para revisar cantidades y responsables, mientras que el mapa reduce la separación entre consulta de stock y búsqueda de ubicación. Las estimaciones y los avisos estadísticos agregan criterios de revisión. No se ha medido una reducción de tiempos o costos en el taller, por lo que esos beneficios permanecen como resultados esperados y no como mejoras cuantificadas.')
h('Objetivos',2)
h('Objetivo general',3)
p('Desarrollar un Sistema Inteligente para la Gestión y Análisis de Inventario en Almacenes mediante Inteligencia Artificial.')
h('Objetivos específicos',3)
for s in [
 'Diseñar un sistema que permita gestionar y controlar las entradas, salidas, existencias y movimientos de los productos dentro del almacén.',
 'Integrar un mecanismo de reconocimiento facial que permita identificar a los empleados autorizados y relacionarlos con los movimientos realizados en el inventario.',
 'Desarrollar un agente inteligente capaz de interpretar consultas en lenguaje natural y proporcionar información sobre los productos, sus existencias, ubicaciones y movimientos del almacén.',
 'Aplicar técnicas de Machine Learning, mediante aprendizaje supervisado y no supervisado, para analizar el comportamiento del inventario, predecir posibles agotamientos de stock y detectar movimientos anormales.'
]: bullet(s)
p('Se conserva la formulación del segundo entregable [2]. En la versión operativa, el término empleados autorizados corresponde al catálogo de empleados activos inscritos para operaciones y reconocimiento, sin control de acceso a la aplicación. La selección del responsable de un movimiento se registra por su clave de empleado; no existe una verificación biométrica obligatoria para cada operación. El vínculo facial automático por movimiento continúa como alcance de integración por evaluar.')
p('La estimación principal de stock y la detección principal de anomalías emplean métodos estadísticos sencillos. Existen componentes heredados de regresión e Isolation Forest, pero su presencia no demuestra que el objetivo de aprendizaje supervisado y no supervisado haya sido evaluado completamente. Su comparación con la solución actual y su validación experimental permanecen pendientes.')
h('Alcance',2)
p('El alcance funcional comprende catálogo de productos, categorías y unidades; existencias por ubicación; entradas, salidas, transferencias y ajustes; historial, lotes y vencimientos; conteos físicos, alertas y reportes. Incluye empleados, pipeline facial, consultas escritas y por voz, mapa lógico y análisis de consumo y cantidades inusuales.')
p('Se excluyen facturación, contabilidad, nómina, robots y automatización física completa. Los avisos, estimaciones y reconocimientos apoyan la revisión humana. No se incluyen módulos de cuentas, permisos o autenticación externa. La señalización con LEDs y ESP32 se mantiene como posibilidad futura y no como funcionalidad del sistema actual.')
h('Recursos e interesados',2)
table('Participación de los interesados',['Interesado','Responsabilidad en el proyecto'],[
 ('Responsable del almacén','Revisar stock, ubicaciones, diferencias, alertas y recomendaciones.'),
 ('Empleados del taller','Registrar o respaldar operaciones y participar en captura y validación facial.'),
 ('Personal de supervisión','Examinar historial y hallazgos de cantidades inusuales.'),
 ('Responsable técnico','Mantener configuración, modelos, conexión SQL y equipos.'),
 ('Equipo desarrollador','Analizar, implementar, documentar y evaluar los componentes académicos.')],[4.5,11.42])
p('Los recursos tecnológicos incluyen una computadora, SQL Server, Python y las bibliotecas instaladas; cámara para imágenes, micrófono y salida de audio para voz. Se requieren imágenes faciales consentidas y representativas e historia operativa suficiente. El equipo académico conserva los integrantes de la portada. No se dispone de un presupuesto verificado de adquisición ni de costos de operación para cuantificar recursos financieros.')
h('Requisitos y restricciones',2)
p('Las operaciones deben registrar producto, cantidad válida, ubicación y empleado activo; la salida no puede superar la disponibilidad y una transferencia debe conservar el stock total. Las consultas deben utilizar datos registrados y mostrar insuficiencia cuando corresponde. La identificación facial debe admitir desconocidos y rechazar imágenes de baja calidad. Los reportes deben conservar el mismo conjunto filtrado de su vista previa.')
p('Las restricciones comprenden disponibilidad de SQL Server local, controladores y dependencias, resolución e iluminación de cámara, condiciones acústicas y volumen de historia. Edge TTS requiere conectividad para generar audio, con alternativa local pyttsx3. La calidad analítica depende de registros completos y del cambio de patrones de consumo. El prototipo no sustituye el conteo físico ni la revisión del responsable.')
h('Riesgos',2)
table('Riesgos y tratamiento',['Riesgo','Consecuencia','Tratamiento'],[
 ('Historia escasa o simulada','Estimaciones poco representativas.','Recopilar historia observada y separar demostración de evaluación.'),
 ('Movimiento incorrecto','Diferencias de stock y análisis distorsionado.','Validar cantidades, usar transacciones y contrastar conteos físicos.'),
 ('Muestras faciales deficientes','Confusión de identidad o rechazo.','Calidad, múltiples muestras, calibración y desconocidos independientes.'),
 ('Cantidad legítima excepcional','Aviso estadístico que no implica irregularidad.','Revisión humana con motivo y estado de justificación.'),
 ('Equipo o audio incompatible','Cámara o voz no disponible.','Comprobaciones explícitas y estados entendibles.'),
 ('Cambios de demanda','Cobertura estimada desactualizada.','Revisar ventana, entradas recientes y plazo real de proveedores.')],[3.9,5.9,6.12])
h('Criterios de éxito',2)
p('El éxito funcional requiere conservar la coherencia entre existencias y movimientos, localizar todas las ubicaciones de un producto, consultar responsables y exportar los resultados filtrados. En JARVIS se debe comprobar que cantidades y ubicaciones coinciden con los servicios y que los nombres ambiguos generan una aclaración. Estos criterios pueden verificarse con casos controlados sin adjudicar porcentajes de eficacia.')
p('La evaluación facial deberá calcular matriz de confusión, precision, recall y F1 sobre muestras independientes y desconocidos; la predicción deberá contrastarse con consumo posterior y errores definidos; las anomalías requieren revisión de falsos avisos y casos pertinentes. Los antecedentes no establecen metas numéricas consolidadas y la revisión actual no aporta esas métricas. Su definición y medición siguen pendientes.')
h('Objetivos de minería de datos',2)
p('Se busca identificar empleados a partir de imágenes válidas, interpretar consultas sobre las entidades del almacén, estimar cobertura con salidas históricas y señalar cantidades que se alejan del comportamiento previo. Producto, tipo de movimiento, fecha y cantidad permiten construir conjuntos analíticos; las muestras faciales deben mantener su relación con empleados activos.')
p('El análisis vigente no incorpora una clasificación validada de fraude ni un modelo independiente de frecuencia operativa. El número de antecedentes sirve para decidir si existe historia suficiente. Ampliar las variables y comparar métodos supervisados o no supervisados requiere datos observados y criterios de evaluación, sin asumir que un modelo más complejo mejora los resultados.')

h('Fase 2. Comprensión de los datos')
h('Fuentes de datos',2)
p('Las fuentes operativas son SQL Server y los registros mantenidos por los servicios del inventario. Las imágenes y embeddings se conservan en archivos del proyecto, relacionados con empleados mediante rostros_empleados. La configuración central se guarda en JSON. Los entregables académicos aportan problema, objetivos y metodología, mientras que las pruebas y capturas aportan evidencia funcional [1–5].')
p('seed_demo.py cargó un escenario reproducible en la base existente: agrega productos y ubicaciones identificados con DEMO-, movimientos con referencias DEMO-V1- y observaciones [DEMO_V1]. Reutiliza las funciones de inventario dentro de una transacción y conserva registros anteriores. La marca final evita repetir la carga. Es historia preparada para demostrar casos, no consumo observado de los empleados mencionados.')
h('Base de datos',2)
p('InventarioIA2 utiliza SQL Server y pyodbc. La inspección actual identifica 20 tablas de aplicación. Las claves primarias y foráneas relacionan catálogos, movimientos y detalles; las cantidades admiten decimales y las reglas de servicio evitan saldos negativos. Las vistas vw_stock_producto y vw_inventario ofrecen la existencia total calculada para consultas compartidas.')
table('Cantidades actuales consultadas en SQL Server',['Entidad','Registros','Interpretación'],[
 (name,E['registros'][name],desc) for name,desc in [
 ('productos','24 registros: 4 anteriores y 20 del escenario demo.'),('categorias','Catálogo reutilizado y ampliado para el escenario.'),('unidades_medida','Unidades conservadas del catálogo existente.'),('ubicaciones','6 anteriores y 12 de demostración.'),('empleados','Catálogo anterior conservado; no creado por el seed.'),('existencias','Saldos por producto y ubicación.'),('movimientos','Operaciones preparadas para demostración.'),('detalle_movimientos','Detalles de las operaciones.'),('lotes','Lotes del escenario con vigencia y vencimientos.'),('existencias_lotes','Distribución por lote y ubicación.'),('rostros_empleados','Relaciones conservadas; no eventos de reconocimiento.'),('anomalias_inventario','Hallazgos calculados por el motor estadístico.'),('alertas','Incluye registros atendidos y pendientes.'),('conteos_inventario','Conteo preparado y aplicado en la demostración.')]], [4.5,2.1,9.32])
p('Las cantidades corresponden a la consulta del 3 de octubre de 2026 y pueden cambiar con nuevas operaciones. La consulta directa registró 22 alertas pendientes entre las 43 filas; las pantallas pueden aplicar criterios de vigencia y actualizaciones derivadas. Un número de filas no equivale a un resultado experimental ni a un porcentaje de calidad.')
h('Descripción de tablas principales',2)
table('Modelo de entidades y relaciones',['Grupo / tablas','Atributos y relación funcional'],[
 ('productos, categorias, unidades_medida','Código, nombre, descripción, categoría, unidad, precio, mínimo, estado y control de lotes. Producto enlaza catálogos.'),
 ('ubicaciones, existencias','Código, sección, pasillo, estante, nivel, almacén y zona; existencia relaciona producto y ubicación con cantidad.'),
 ('empleados, rostros_empleados','Código, nombres, apellidos, cargo y estado; relación entre empleado y subject facial.'),
 ('movimientos, detalle_movimientos','Fecha, tipo, empleado, motivo y referencia; detalles con producto, cantidad, origen/destino y saldos.'),
 ('tipos_movimiento, motivos_movimiento','Catálogos de operaciones y motivos; contienen 5 y 8 registros, respectivamente.'),
 ('lotes, existencias_lotes, detalle_movimiento_lotes','Lote por producto, fechas y distribución; detalle de operación enlaza sus cantidades por lote.'),
 ('alertas, tipos_alerta','Tipo, mensaje, nivel, entidades relacionadas y atención; 9 tipos disponibles.'),
 ('conteos_inventario, detalle_conteos','Cabecera de conteo y diferencias por producto/ubicación; aplicación mediante ajuste.'),
 ('anomalias_inventario','Detalle de movimiento, tipo, nivel, score, motivo, referencia histórica, estado y observación.'),
 ('inventario_migraciones','Control técnico de migraciones de inventario aplicadas.')],[6.1,9.82])
p('En la consulta existen 8 filas de existencias_lotes, 46 vínculos en detalle_movimiento_lotes, un detalle de conteo y un registro en inventario_migraciones. Las tablas de detalle no deben interpretarse como productos, empleados ni eventos faciales adicionales.')
h('Productos',2)
p('Cada producto tiene código único y una unidad definida. El stock se calcula a partir de existencias; crear un artículo no inserta una cantidad inicial ficticia. Stock mínimo permite distinguir avisos de disponibilidad. Los campos de códigos de barras y QR existen, pero los 24 registros consultados tienen esos campos nulos; su presencia no acredita una integración con lectores físicos.')
p('El escenario contiene tornillería, filtros, lubricantes, herramientas, equipos y consumibles. Entre sus 20 productos hay 13 con stock normal, 4 bajos, 2 críticos y uno agotado según las reglas de preparación. Las cantidades se expresan en Unidad, Litro, Metro, Kg o Paquete. Sumar unidades distintas no produce un stock físico comparable, por lo que los reportes separan totales por unidad.')
h('Empleados',2)
p('Los 16 empleados ya existían antes de cargar la demostración. Las 244 operaciones se asignaron de forma simulada a empleados del catálogo; esa asignación no acredita trabajo realmente realizado por esas personas. El movimiento conserva la clave del responsable y las relaciones faciales conservan el subject utilizado para identificación.')
p('rostros_empleados contiene 16 relaciones, pero no representa un historial persistente de reconocimientos exitosos. La actividad facial reciente de la sesión puede mostrarse en la interfaz. No se fabricaron capturas, embeddings ni eventos para completar reportes; el historial facial persistente continúa sin datos.')
h('Ubicaciones',2)
p('La jerarquía disponible comprende almacén, zona o sección, pasillo, estante y nivel dentro de la tabla ubicaciones. No hay columnas operativas de posición o capacidad en el esquema revisado. Las seis ubicaciones anteriores tienen almacén y zona nulos; el mapa debe presentar la información disponible sin atribuirles valores físicos inventados.')
p('El escenario agrega 12 ubicaciones en Almacén de demostración, distribuidas entre zonas A, B y C. Cinco productos demo ocupan dos ubicaciones. DEMO-TOR-01, por ejemplo, tiene 170 unidades en Zona A/E01 y 70 en Zona B/E03, para un total de 240. La representación del mapa es lógica y no contiene coordenadas de un plano a escala.')
h('Movimientos',2)
table('Historia disponible por tipo de operación',['Tipo','Cantidad','Efecto sobre el inventario'],[
 ('Entrada',26,'Incrementa stock en destino.'),('Salida',211,'Reduce stock en origen y constituye consumo para predicción.'),('Transferencia',6,'Redistribuye stock entre ubicaciones; no es consumo global.'),('Ajuste salida',1,'Regulariza una diferencia de conteo; se excluye del consumo.'),('Total',244,'Historia del escenario de demostración.')],[4.1,2.2,9.62])
p('La primera operación se registró el 10 de julio de 2026 a las 08:00 y la última el 2 de octubre de 2026 a las 17:00. Los detalles conservan cantidades y ubicaciones. El conteo preparado de brocas comparó 8 unidades registradas con 6 físicas y produjo un ajuste de salida de 2; se trata de un caso simulado de conciliación.')
h('Calidad de los datos',2)
table('Comprobaciones de calidad realizadas en SQL',['Comprobación','Resultado','Alcance'],[
 ('Existencias negativas',0,'Cantidad menor que cero en existencias.'),('Códigos de producto duplicados',0,'Grupos repetidos en productos.codigo.'),('Pares producto–ubicación duplicados',0,'Grupos repetidos en existencias.'),('Cantidades no positivas',0,'Detalles con cantidad menor o igual a cero.'),('Movimientos futuros',0,'Fecha posterior al reloj de SQL Server.'),('Movimientos sin empleado',0,'Clave id_empleado nula.'),('Referencias huérfanas',sum(x['huerfanos'] for x in E['integridad_foranea']),'Cruces de las claves foráneas inspeccionadas.'),('Códigos de barras y QR nulos','24 en cada campo','Datos opcionales aún sin completar.'),('Almacén y zona nulos','6 en cada campo','Ubicaciones anteriores con jerarquía parcial.')],[6.2,2.3,7.42])
p('Estas verificaciones describen el contenido actual y no garantizan que los registros correspondan a la realidad física. La consistencia del escenario también se comprobó reconstruyendo saldos desde movimientos, con cero diferencias y sin saldos negativos en la secuencia. Para la operación real se requiere cotejar ubicaciones, unidades, lotes y cantidades mediante procedimientos del almacén.')
p('Los nulos opcionales deben interpretarse por su función: un vencimiento ausente puede corresponder a un lote sin fecha definida; una ubicación sin almacén o zona limita el detalle del mapa. No procede reemplazar esos valores por descripciones simuladas ni usar la ausencia como evidencia de calidad facial. La calibración de embeddings es un requisito independiente.')
h('Datos disponibles para análisis',2)
p('La predicción utiliza fechas y cantidades de salidas agregadas por producto. La demostración permite que 12 productos tengan historia suficiente con la configuración verificada; los demás conservan la respuesta de insuficiencia. Ese resultado depende de la ventana y de la fecha de análisis, y no valida el error de predicción frente a demanda futura real.')
p('La detección de anomalías compara cada cantidad con antecedentes previos del mismo producto y tipo, excluyendo el movimiento candidato de su referencia. Los tres hallazgos del escenario proceden de salidas deliberadamente excepcionales. Las imágenes y la galería permiten preparar la identificación facial, pero no existe en la revisión una calibración válida guardada para confirmar identidades con el motor principal.')

h('Estado actual del sistema')
h('Arquitectura',2)
p('main.py inicia interfaz.py. Las pantallas solicitan operaciones a módulos de dominio y a inventario_service; database.py abre conexiones SQL mediante pyodbc. Los servicios reutilizan entidades y reglas de stock. JARVIS selecciona herramientas de consulta y devuelve texto, datos y acciones de interfaz para abrir o resaltar el mapa. Configuración y estilos tienen módulos centrales.')
figure(diagram('arquitectura',['Operador','Tkinter /\nttkbootstrap','Servicios de\ninventario y análisis','pyodbc','SQL Server']), 'Arquitectura de la aplicación de escritorio.')
p('La cámara, el audio y varias consultas se procesan en tareas independientes; las colas llevan resultados al hilo de Tkinter para actualizar widgets. Los archivos de modelos, imágenes y JSON complementan los datos SQL. La separación permite que un error de cámara o micrófono se informe sin fabricar resultados del inventario.')
h('Tecnologías',2)
table('Tecnologías comprobadas en el código',['Tecnología','Uso actual'],[
 ('Python; Tkinter; ttkbootstrap','Aplicación de escritorio y tema flatly.'),('SQL Server; pyodbc','Persistencia relacional y acceso a datos.'),('OpenCV; YuNet; SFace','Captura, detección, alineación y embeddings del motor facial principal.'),('NumPy; Pandas; scikit-learn','Vectores, series, análisis y componentes de regresión/Isolation Forest.'),('TensorFlow / Keras','CNN facial heredada como alternativa separada.'),('Vosk; sounddevice','Modelo de reconocimiento de voz local y captura de micrófono.'),('edge-tts; pyttsx3','Síntesis de voz con alternativa local.'),('Matplotlib','Visualización de consumo y datos analíticos.'),('ReportLab; openpyxl','Exportación PDF y Excel.'),('JSON; pathlib; threading; queue','Configuración, archivos y ejecución de tareas.')],[5.6,10.32])
p('Existe un cliente opcional para interpretación estructurada mediante Ollama local, con validación de acciones y fallback a reglas. Su presencia en llm_jarvis.py no demuestra que el servicio o el modelo estén activos. JARVIS conserva consultas deterministas y respuestas derivadas de datos; no se presenta como un modelo generativo general.')
h('Base de datos y servicios',2)
p('Las migraciones existentes de inventario y anomalías extienden el esquema para lotes, conteos y hallazgos. No se crean tablas duplicadas para las pantallas o JARVIS. Los movimientos se ejecutan con transacciones y bloqueos de stock por producto, validación de entidades activas y cantidades decimales. Una excepción revierte la operación incompleta.')
p('El stock total se obtiene desde las vistas y servicios comunes, no desde un valor independiente escrito por cada módulo. Un lote relaciona producto y cantidades por ubicación; sus movimientos mantienen vínculos con el detalle correspondiente. El precio de producto no se interpreta automáticamente como costo de compra ni como valoración contable certificada.')
h('Interfaz gráfica',2)
p('La interfaz utiliza navegación lateral, tarjetas, tablas, filtros, estados y formularios. ui_theme.py centraliza colores, tipografía, espaciado y estilos para mantener consistencia visual. Las 15 páginas siguen conectadas con los mismos servicios; los estados vacíos informan ausencia de datos o insuficiencia en lugar de mostrar valores preparados en la pantalla.')
p('Las capturas del anexo corresponden a la interfaz revisada y al escenario demo. Ilustran navegación y presentación, no resultados biométricos ni rendimiento de hardware. Los contadores y estados pueden variar con la fecha, las consultas y nuevas operaciones. Se incluyen vistas de inventario, movimientos, mapa, empleados, análisis, anomalías, reportes, JARVIS y configuración.')

h('Módulos implementados')
modules=[
 ('Dashboard','IMPLEMENTADA. Resume productos activos, stock, productos bajos y agotados, alertas, empleados, movimientos y anomalías. Incluye accesos rápidos, últimos movimientos, productos críticos, alertas y actividad facial reciente de sesión. Los gráficos se construyen con datos disponibles. Los estados de SQL, reconocimiento y JARVIS no sustituyen pruebas físicas.','La suma de cantidades de distintas unidades debe leerse como indicador agregado de registros, no como una magnitud física homogénea; para comparar stock se utilizan unidades y detalle de producto.'),
 ('Productos','IMPLEMENTADA. Registra, consulta, edita y desactiva productos, con código, nombre, categoría, unidad, precio, mínimo y control de lotes. Permite filtros, fichas y Ver ubicación para abrir el mapa con el producto seleccionado. Crear un producto no aumenta stock: su cantidad se incorpora por entrada.','Las unidades y cantidades decimales se conservan durante operaciones. Campos opcionales como códigos de barras, QR o imagen no implican que exista un dispositivo lector integrado.'),
 ('Movimientos','IMPLEMENTADA. Registra entradas, salidas, transferencias y ajustes con empleado, producto, cantidad, origen/destino, motivo, referencia y observación. La validación impide cantidades inválidas, uso de entidades inactivas y salidas superiores al saldo. Las operaciones se confirman de forma transaccional.','Los lotes se seleccionan según disponibilidad y vigencia, con sugerencias FEFO/FIFO. Una salida de lote vencido requiere confirmación explícita. Las transferencias conservan stock global y los costos opcionales de entrada no reemplazan el precio del catálogo.'),
 ('Historial','IMPLEMENTADA. Presenta las operaciones registradas y permite filtrar y consultar sus detalles, fechas, cantidades, empleados y ubicaciones. La relación entre cabecera y detalle facilita revisar cómo cambió una existencia.','Este historial es trazabilidad operativa del inventario. No es una bitácora de auditoría de usuarios ni certifica que el responsable fue reconocido biométricamente en cada movimiento.'),
 ('Alertas','IMPLEMENTADA. Muestra avisos derivados del stock, vencimientos y cantidades inusuales, junto con avisos operativos existentes. Permite revisar y atender avisos; los servicios actualizan estados cuando cambian las condiciones.','Un aviso de agotamiento estimado o cantidad inusual necesita interpretación humana. Una alerta no demuestra fraude, error definitivo ni demanda futura cierta.'),
 ('Empleados','IMPLEMENTADA la administración operativa. Conserva código, nombres, apellidos, cargo y estado, y permite relacionar imágenes o subject con un empleado. El panel facial ofrece elección de motor, registro, regeneración y reconocimiento con cámara.','EN PROCESO la calibración del motor principal. La identidad confirmada requiere política válida y empleado activo. La sección no crea cuentas de acceso; el reporte facial no tiene reconocimientos persistentes fabricados.'),
 ('Mapa del almacén','IMPLEMENTADA. Genera agrupaciones desde ubicaciones SQL: almacén, zona o sección, pasillo, estante y nivel. Permite filtrar, buscar producto o ubicación, seleccionar y consultar productos con cantidad y unidad. Resalta todas las ubicaciones de un producto.','mostrar_producto_en_mapa integra Productos y JARVIS; las operaciones pueden consultar origen y destino. La vista usa textos y estados junto con colores. Es un mapa lógico: no hay posición, capacidad ni coordenadas físicas para calcular llenado o dibujar un plano a escala.'),
 ('Conteo de inventario','IMPLEMENTADA. Compara cantidades físicas con las registradas por ubicación y conserva borradores. Al aplicar un conteo, la diferencia se convierte en ajuste relacionado con su detalle. La validación detecta cambios del saldo esperado y evita aplicar dos veces el mismo conteo.','El caso demo de brocas permite verificar la conciliación y el historial. No acredita que se haya contado físicamente el almacén ni sustituye un procedimiento de conteo real.'),
 ('Análisis IA','IMPLEMENTADA. Mantiene consultas de consumo, gráficas y acceso a predicción y detección de anomalías. La serie histórica proviene de salidas registradas y los resultados utilizan los motores compartidos.','Los resultados deben leerse con su ventana y suficiencia. Una gráfica de historia preparada no demuestra la generalización de un modelo sobre demanda del taller.'),
 ('JARVIS','IMPLEMENTADA. Atiende texto escrito y voz transcrita mediante un mismo procesador. Consulta productos, stock, mínimos, agotados, ubicaciones, movimientos, empleados, predicción y anomalías; puede abrir el mapa y resaltar resultados.','Su capa conversacional conserva contexto temporal, formula respuestas breves, aclara productos similares y comunica errores. Push-to-talk es principal; activación por JARVIS es opcional. No ejecuta entradas ni salidas por voz.'),
 ('Predicción de stock','IMPLEMENTADA. Muestra stock actual, consumo diario observado, tendencia, cobertura, fecha aproximada y riesgo cuando existe historia suficiente. Ofrece tabla general y gráfica, con ventanas configurables.','El método principal calcula stock / consumo promedio. Entradas, transferencias y ajustes no son consumo. La cantidad sugerida deriva de cobertura objetivo y mínimo; su utilidad operativa y error requieren validación con datos futuros reales.'),
 ('Anomalías','IMPLEMENTADA. Analiza cantidades frente a historia previa por producto y tipo; presenta fecha, cantidad, score, nivel, motivo y referencia histórica. Los hallazgos se guardan sin duplicar el mismo detalle y tipo.','Permite estados PENDIENTE, REVISADA y JUSTIFICADA; justificar exige observación. La pantalla, los reportes y JARVIS comparten los hallazgos. La revisión no equivale a auditoría ni a una clasificación de fraude.'),
 ('Reportes','IMPLEMENTADA. Ofrece vista previa y filtros por las entidades o fechas pertinentes. Genera PDF con ReportLab y Excel con openpyxl desde el mismo conjunto de datos mostrado, conservando filtros, encabezado y fecha.','Hay 15 tipos: inventario actual, stock bajo, agotados, entradas, salidas, transferencias, historial, inventario por ubicación, productos por categoría, predicción, anomalías, alertas, empleados, reconocimiento facial y resumen general. Excel conserva números/fechas; los totales de stock se separan por unidad.'),
 ('Estado del sistema','IMPLEMENTADA. Comprueba conexión SQL y componentes de JARVIS, modelos faciales, voz, predicción y anomalías. Incluye pruebas explícitas de cámara, micrófono y reproducción, con resultado y momento de comprobación.','Son diagnósticos puntuales, no monitoreo permanente. Un modelo existente no garantiza calibración; un archivo Vosk no equivale a reconocimiento acústico validado y un equipo sin prueba debe conservar su estado sin verificar.'),
 ('Configuración','IMPLEMENTADA. config_manager.py valida y guarda de forma atómica config_sistema.json. Centraliza opciones generales, inventario, fuente y resolución de cámara, calidad facial, micrófono, voz, activación, análisis, predicción, anomalías y reportes.','La configuración admite ventanas de 7, 30, 60 y 90 días y valida rangos. Mantiene compatibilidad con configuración facial previa. El umbral de identificación depende de una calibración coherente, no de un ajuste libre para forzar reconocimientos.')
]
for title,a,b in modules:
    h(title,2); p(a); p(b)

h('Componentes de Inteligencia Artificial')
h('Reconocimiento facial',2)
p('El pipeline principal usa YuNet para detección y landmarks, y SFace para embeddings mediante OpenCV. Los fotogramas llegan en BGR; el recorte se alinea con alignCrop y el vector se normaliza. Registro y reconocimiento comparten esa extracción para evitar comparar imágenes procesadas de manera incompatible. La galería conserva múltiples muestras y su regeneración aprovecha imágenes existentes, sin mezclar embeddings de modelos distintos.')
figure(diagram('facial',['Cámara\nBGR','YuNet +\ncalidad','Landmarks y\nalineación','SFace +\nnormalización','Comparación +\ntracking','Empleado o\ndesconocido']), 'Pipeline facial principal; confirmación sujeta a calibración válida.')
p('Antes de comparar se revisan tamaño en píxeles, nitidez, iluminación, orientación y visibilidad. El seguimiento temporal acumula evidencia y consenso; un fotograma deficiente puede producir VERIFICANDO o BAJA CALIDAD en vez de una identidad forzada. La decisión considera mejor candidato, segundo candidato, margen y umbral calibrado. La similitud no es un porcentaje de certeza y no debe interpretarse como accuracy.')
p('La revisión no encontró calibracion.json en la galería activa. La infraestructura está implementada, pero la confirmación de identidades del motor principal permanece condicionada a generar y validar una política vinculada a esa galería, modelos y configuración. No hay evidencia para publicar precision, recall, F1 o alcance en metros. CERCA, MEDIA y LEJOS representan tamaño aparente de cara, no distancia física medida.')
p('La fuente de vídeo está separada del algoritmo y admite índice de webcam/USB o fuentes IP/RTSP. La captura y el procesamiento evitan acumular fotogramas obsoletos mediante tareas y colas. La capacidad efectiva depende de resolución, lente, iluminación y píxeles disponibles en el rostro. La CNN TensorFlow/Keras permanece como alternativa separada; el código antiguo Haar/LBPH no define el pipeline principal documentado.')
h('Procesamiento de lenguaje natural y JARVIS',2)
p('El agente normaliza texto, identifica intenciones y entidades y selecciona herramientas del inventario. Productos similares generan opciones antes de adoptar una identidad de producto. El contexto temporal permite continuar una consulta con ¿Dónde está? o ¿Cuánto durará? después de identificar el artículo. jarvis_personality.py aporta respuestas naturales y mensajes breves sin inventar cantidades.')
figure(diagram('jarvis',['Micrófono','sounddevice\n+ Vosk','Procesador\ncompartido','Servicios +\nSQL','Texto y\nacciones UI','Edge TTS /\npyttsx3']), 'Flujo de voz de JARVIS; el texto escrito entra al mismo procesador.')
p('Vosk realiza reconocimiento local en español y sounddevice recibe audio. La transcripción aparece como mensaje del operador y se procesa igual que una consulta escrita. Edge TTS produce voz con conectividad y pyttsx3 sirve como alternativa local. Escucha y reproducción se coordinan para evitar eco y las tareas de audio no se ejecutan como esperas prolongadas en el hilo principal de Tkinter.')
p('JARVIS comunica silencio, micrófono ocupado o inexistente, modelo Vosk ausente y errores de síntesis. Puede consultar relaciones faciales registradas y compartir estados, sin sustituir la detección de cámara. El cliente opcional Ollama interpreta acciones estructuradas permitidas, no produce SQL libre ni cifras que reemplacen resultados de servicios. Su disponibilidad efectiva requiere una comprobación adicional.')
h('Predicción de stock',2)
p('La técnica principal es el promedio diario observado dentro de una ventana. Se agregan las salidas por fecha y se incluyen días sin consumo desde el comienzo de la serie disponible hasta la fecha de análisis. La configuración permite 7, 30, 60 o 90 días; el mínimo vigente de suficiencia es 5 días distintos con salidas. Si no se cumple, se devuelve Datos insuficientes; si no hay consumo, no se asigna una fecha falsa.')
table('Cálculos del motor de predicción',['Indicador','Regla implementada'],[
 ('Consumo promedio diario','Suma de salidas / días de la serie analizada, incluyendo días sin salida.'),
 ('Cobertura estimada','Stock total del producto / consumo promedio positivo.'),
 ('Fecha aproximada','Fecha de análisis + cobertura redondeada hacia arriba; exige stock y consumo positivos.'),
 ('Tendencia de consumo','Comparación de promedio reciente y anterior; cambios mayores que el margen del código determinan aumento o disminución.'),
 ('Stock objetivo','Máximo entre stock mínimo y consumo promedio × horizonte de cobertura.'),
 ('Cantidad sugerida','Máximo entre stock objetivo − stock actual y cero.'),
 ('Riesgo','Combina mínimo y cobertura: CRÍTICO, ALTO, MEDIO o BAJO; SIN DATOS cuando falta historia.')],[4.5,11.42])
p('La regla central asigna CRÍTICO ante stock no positivo, stock igual o menor a la mitad del mínimo o cobertura de hasta 3 días. ALTO considera stock igual o menor al mínimo o cobertura de hasta 7 días; MEDIO corresponde a cobertura de hasta 30 días, y BAJO al resto. Estos límites son reglas del código, no umbrales de eficacia obtenidos experimentalmente.')
p('La estimación supone continuidad del consumo y ausencia de entradas futuras. No incorpora plazos de proveedor verificados, estacionalidad ni demanda extraordinaria. Una regresión lineal se conserva como cálculo complementario, pero la cobertura principal procede del promedio. Las salidas excepcionales influyen en ese promedio y requieren interpretación. La exactitud de la fecha de agotamiento está PENDIENTE DE VALIDACIÓN.')
h('Detección de anomalías',2)
p('El motor principal ordena movimientos cronológicamente y agrupa antecedentes por producto y tipo. Cada candidato se compara solo con registros anteriores, dentro del período configurado; el mínimo vigente es 10 antecedentes y la sensibilidad predeterminada es 3. Cuando existe variación, el score es (cantidad − promedio previo) / desviación estándar poblacional. Si la variación es nula, se usa la razón cantidad / promedio positivo.')
p('Se señala una cantidad cuando supera el promedio y alcanza la sensibilidad. Los niveles del código distinguen CRÍTICA desde score 6, ALTA desde 4 y MEDIA desde 3; sensibilidades menores pueden producir nivel BAJA. El motivo informa promedio, desviación, número de muestras y métrica. El score expresa separación respecto al historial, no probabilidad de fraude ni severidad física garantizada.')
p('Las variables centrales son cantidad, producto, tipo y orden temporal. La frecuencia aporta disponibilidad de antecedentes, pero no constituye actualmente un detector independiente de patrones horarios o frecuencia de retiros. Isolation Forest se conserva como alternativa heredada con cantidad, hora, día y stock anterior; no es el motor principal compartido por la pantalla actual y los reportes. La pertinencia de los hallazgos necesita revisión del almacén.')

h('Integración entre módulos')
p('Productos, Movimientos y Mapa utilizan las mismas ubicaciones y existencias. Una búsqueda desde Productos abre el mapa con todas las ubicaciones del artículo; una transferencia consulta origen y destino. El stock global sigue siendo la suma de existencias y no se reduce por mover material dentro del almacén. El conteo físico genera ajustes identificados, manteniendo la diferencia separada del consumo histórico.')
figure(diagram('integracion',['Existencias\npor ubicación','Movimientos\ny salidas','Predicción /\nAnomalías','Pantallas y\nreportes','JARVIS +\nmapa']), 'Datos operativos y motores analíticos compartidos entre módulos.')
p('JARVIS consulta esos servicios y devuelve acciones de interfaz para localizar productos. Una consulta exacta de DEMO-TOR-01 obtiene 240 unidades; la localización identifica 170 en Zona A/E01 y 70 en Zona B/E03. Una pregunta posterior sobre duración utiliza predecir_agotamiento, con su ventana y suficiencia. Si se pregunta solo por tornillos o filtros y existen varias coincidencias, se solicita aclaración antes de responder por un producto específico.')
p('Las anomalías guardadas se comparten con JARVIS y Reportes; las predicciones se recalculan con el motor central y no con fórmulas duplicadas en el asistente. Vista previa y exportación conservan el conjunto filtrado. Empleados enlaza responsables y subjects faciales, mientras que la integración biométrica obligatoria de cada movimiento no se considera implementada.')

h('Resultados actuales')
p('La versión dispone de las 15 páginas operativas descritas. SQL Server contiene el escenario de demostración y mantiene el catálogo anterior. La verificación actual confirmó sus cantidades, fechas y relaciones sin modificar el contenido de la base. Las salidas inusuales preparadas originaron tres hallazgos calculados, y la historia disponible permite mostrar tanto predicciones con suficiencia como respuestas sin datos.')
table('Evidencia funcional disponible',['Resultado','Evidencia y límite'],[
 ('Consistencia del escenario','244 movimientos y 244 detalles; reconstrucción de stock sin diferencias según verificación existente.'),
 ('Idempotencia del seed','Una segunda ejecución desde conexión nueva agregó cero registros.'),
 ('Pruebas de regresión','86 pruebas existentes aprobadas tras la carga; verifican comportamiento, no accuracy de IA.'),
 ('Revisión de interfaz','45 comprobaciones de pantallas en tres tamaños; sin errores de callbacks ni botones recortados en esa revisión.'),
 ('Integración de consultas','JARVIS consultó datos SQL y presentó aclaración ante productos similares.'),
 ('Anomalías del escenario','3 salidas deliberadamente inusuales detectadas y explicadas; no son casos reales de irregularidad.'),
 ('Reconocimiento facial','Pipeline y galería disponibles; calibración válida y métricas físicas pendientes.')],[5.5,10.42])
p('Las salidas excepcionales corresponden al tornillo hexagonal (150 unidades el 28 de septiembre), filtro de aceite (25 unidades el 29 de septiembre) y cable eléctrico (80 metros el 30 de septiembre). Esas cantidades fueron preparadas en el seed; los scores y motivos se calcularon con el motor. El resultado demuestra ejecución del flujo estadístico en un caso controlado, sin estimar sensibilidad, especificidad ni falsos positivos en operación real.')
p('Las capturas muestran el sistema actual con tablas y filtros poblados. No existe un ensayo publicado de rendimiento temporal ni comparación independiente de modelos. Tampoco se presentan resultados de reconocimiento en metros, reducción de costos, tiempos de búsqueda o aceptación del personal, porque no se han medido en los archivos de evidencia revisados.')

h('Limitaciones actuales')
table('Limitaciones y estado de validación',['Componente','Limitación','Estado'],[
 ('Datos históricos','La historia principal de movimientos fue preparada para demostración.','PENDIENTE DE VALIDACIÓN con registros observados.'),
 ('Facial','Calibración vinculada a la galería ausente; sin evaluación independiente.','EN PROCESO calibración; precisión pendiente.'),
 ('Cámara y audio','Sin ensayo documentado de campo en iluminación y ruido del taller.','PENDIENTE DE VALIDACIÓN física.'),
 ('Mapa','Jerarquía parcial en ubicaciones anteriores; sin capacidad, posición o escala.','IMPLEMENTADA representación lógica.'),
 ('Predicción','Promedio sensible a patrones cambiantes; sin evaluación de error futura.','IMPLEMENTADA estimación; utilidad pendiente.'),
 ('Anomalías','Detección centrada en cantidades, sin etiquetas verificadas de irregularidad.','IMPLEMENTADA; ampliación y pertinencia pendientes.'),
 ('Historial facial','No hay tabla de eventos exitosos persistentes; actividad de sesión limitada.','Pendiente definir almacenamiento si se requiere.'),
 ('LEDs / ESP32','No se encontró integración funcional de hardware.','TRABAJO FUTURO.')],[3.5,8.4,4.02])
p('La identificación de empleados registrados no prueba automáticamente quién realizó una operación. La responsabilidad se conserva por el empleado seleccionado en el movimiento. La vinculación biométrica obligatoria requiere definir el flujo y validarlo sin convertirlo en un módulo de autenticación. Los informes y diagnósticos son herramientas de consulta; su existencia no demuestra despliegue continuo ni precisión del análisis.')

h('Conclusiones')
p('El proyecto ha concretado una base operativa centralizada para el inventario del taller: productos, existencias por ubicación, movimientos y responsables se mantienen relacionados en SQL Server. Las entradas, salidas, transferencias, lotes y conteos utilizan validaciones compartidas y permiten revisar el origen de un cambio de cantidad. La demostración conserva información previa y proporciona casos coherentes para recorrer la aplicación.')
p('El mapa lógico integra ubicación y disponibilidad, incluyendo productos distribuidos en más de un lugar. JARVIS conecta las consultas escritas y por voz con los mismos servicios y motores del sistema; su contexto temporal y aclaraciones permiten continuar consultas sin asumir un producto ambiguo. Esta integración está respaldada por el código y comprobaciones funcionales, aunque su utilidad cotidiana requiere observar el uso en el taller.')
p('El reconocimiento facial dispone de detección YuNet, embeddings SFace, alineación, calidad y seguimiento, con una alternativa CNN separada. La calibración y las pruebas físicas siguen pendientes. La predicción produce estimaciones explicadas a partir de salidas y las anomalías muestran cantidades apartadas del historial; ninguno de esos resultados se interpreta como certeza futura, fraude detectado o éxito experimental validado.')
p('CRISP-DM conserva la relación entre necesidades del negocio, datos disponibles y técnicas aplicadas. El siguiente avance debe concentrarse en recopilar información observada, calibrar la visión y evaluar los resultados con criterios definidos. La funcionalidad actual permite esa evaluación, pero no sustituye la evidencia necesaria para cerrar las fases de modelado, evaluación y despliegue en condiciones reales.')

h('Recomendaciones y trabajo futuro')
p('Se recomienda comenzar por validar el catálogo físico: revisar unidades, códigos, almacén, zona, pasillo, estante y nivel; completar campos solo con información confirmada y efectuar conteos por ubicación. La historia observada deberá separarse claramente del escenario demo para evaluar consumo, frecuencia y reabastecimiento sin mezclar operaciones simuladas con actividad del personal.')
p('Para reconocimiento facial, se debe completar la calibración con muestras conocidas y desconocidas independientes, mantener el mismo pipeline y conservar las versiones de modelos y galería. Las pruebas deben registrar iluminación, orientación y tamaño del rostro en píxeles, además de cámara y lente. Se requiere medir errores de identificación y rechazo antes de publicar métricas o rangos físicos de alcance.')
p('La predicción debe compararse con consumo posterior mediante cortes temporales, errores definidos y ventanas equivalentes. El promedio observado constituye una referencia sencilla para contrastar regresión u otros métodos. Conviene incorporar plazos de suministro únicamente cuando existan datos confiables, y revisar cómo las cantidades excepcionales afectan la recomendación de cobertura.')
p('Las anomalías deben revisarse con el responsable del almacén para distinguir reabastecimientos legítimos, retiros extraordinarios y errores de captura. La ampliación a frecuencia u horarios requiere atributos y criterios de negocio suficientes. Cualquier comparación con Isolation Forest debe evitar usar el mismo conjunto para ajustar y evaluar, y documentar falsos avisos.')
p('Se recomienda probar micrófono, síntesis y cámara en el entorno de operación y registrar resultados de diagnóstico puntuales. La evaluación de JARVIS debe incluir nombres similares, consultas combinadas, referencias de contexto y errores de conexión. La voz continuará como medio de consulta mientras no se diseñe y valide un flujo específico de registro de movimientos.')
p('Una integración de señalización con LEDs y ESP32 podría relacionar producto, ubicación y dispositivo, utilizando las acciones existentes de mapa y JARVIS. Antes de incorporarla deberán definirse asociación física, comunicación, estados y pruebas del hardware. No se atribuye a la versión actual control de LEDs ni conexión operativa con ESP32.')

D.add_page_break()
x=D.add_paragraph('Referencias y evidencia del proyecto',style='Title')
for s in [
 '[1] Equipo del proyecto. Primer Entregable IA 1.0. Universidad Nacional de Ingeniería, 4 de septiembre de 2026. Antecedente académico y portada institucional.',
 '[2] Equipo del proyecto. Borrador del segundo entregable (1) (1). Fases de comprensión del negocio y comprensión de los datos. Versión académica de referencia para objetivos y CRISP-DM.',
 '[3] Código de la aplicación revisado al 3 de octubre de 2026: main.py, interfaz.py, database.py, productos.py, movimientos.py, empleados.py, alertas.py, inventario_service.py y módulos de mapa, conteo, reportes, configuración y estado.',
 '[4] Componentes analíticos y de interacción: agente_inventario.py, asistente_voz.py, nlp_jarvis.py, jarvis_personality.py, llm_jarvis.py, prediccion_stock.py, deteccion_anomalias.py, pipeline_facial.py, reconocimiento_embeddings.py y CNN heredada.',
 '[5] SQL Server / InventarioIA2: inspección de tablas, claves, registros, fechas, nulos y consistencia del 3 de octubre de 2026. Documentación del escenario: seed_demo.py, DATOS_DEMO.md y verificar_datos_demo.sql.',
 '[6] Evidencia funcional del proyecto: resultados de pruebas y capturas actuales de verificacion_ui; informe funcional y documentación de integración. Las comprobaciones de interfaz no constituyen experimentos de precisión analítica.'
]: p(s)

D.add_page_break()
x=D.add_paragraph('Anexo. Capturas de la interfaz actual',style='Title')
p('Vistas del prototipo con datos de demostración. Los estados y contadores corresponden al momento de cada captura y pueden cambiar al actualizar consultas. La vista de Empleados ilustra controles faciales; no acredita reconocimiento exitoso ni alcance físico.')
shots=[(0,'Dashboard: indicadores y accesos rápidos.'),(1,'Productos: catálogo y disponibilidad.'),(2,'Movimientos: operación y consultas.'),(8,'Mapa del almacén: jerarquía lógica obtenida de SQL.'),(5,'Empleados: registros operativos y controles faciales.'),(6,'Análisis IA: acceso a consumo y análisis.'),(11,'Anomalías: cantidades inusuales y motivo.'),(12,'Reportes: tipos de reporte y filtros.'),(7,'JARVIS: consulta escrita y controles de voz.'),(14,'Configuración: opciones centrales del sistema.')]
for i,(idx,title) in enumerate(shots):
    if i>0 and i%2==0: D.add_page_break()
    figure(ROOT/'verificacion_ui'/f'modulo_{idx}.png',title)

OUT=Path(sys.argv[1]) if len(sys.argv)>1 else ROOT/'Sistema_Inteligente_Inventario_Actualizado.docx'
D.save(OUT)
(BASE/'resumen_documento.json').write_text(json.dumps({'tablas':T,'figuras':F,'diagramas':4,'capturas':10,'parrafos':len(D.paragraphs),'palabras':sum(len(x.text.split()) for x in D.paragraphs)},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUT); print('Tablas',T,'Figuras',F,'Palabras',sum(len(x.text.split()) for x in D.paragraphs))
