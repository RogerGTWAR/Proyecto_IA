from core.formato_numerico import format_numero, format_moneda
'Renderizadores sin acceso a SQL: reciben la instantánea de integrations/reporting/reportes.py.'
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import math
from xml.sax.saxutils import escape
from core.config_manager import obtener
from integrations.reporting.report_theme import PALETA, FORMATOS, NUMERICAS, FUENTE_EXCEL, colores_estado, ancho_columna

def _datos(r):
    from integrations.reporting.reportes import preparar_reporte
    return r if 'indicadores' in r else preparar_reporte(r)

def _filtros(r):
    from integrations.reporting.reportes import etiqueta
    return ', '.join(f'{etiqueta(k)}: {r.get("filtros_etiquetas",{}).get(k,v)}' for k,v in r['filtros'].items()) or 'Todos'

def _valor(k,v):
    from integrations.reporting.reportes import texto
    if v is None:
        if k=='fecha_agotamiento': return 'Sin fecha estimada'
        return 'Datos insuficientes' if k=='dias_restantes_promedio' else 'Sin datos'
    if k in ('precio','valor') or k=='Valor a precio registrado': return format_moneda(v)
    if k in NUMERICAS: return format_numero(v,agrupar=True)
    return texto(v)

def _guardar(contenido,archivo):
    """La creación exclusiva evita sobrescribir incluso si cambia el destino durante la exportación."""
    from integrations.reporting.reportes import ruta_disponible
    archivo=Path(archivo); archivo.parent.mkdir(parents=True,exist_ok=True)
    while True:
        destino=ruta_disponible(archivo)
        try:
            with destino.open('xb') as f: f.write(contenido)
            return destino
        except FileExistsError: continue

def exportar_pdf(reporte,archivo):
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,LongTable,Table,TableStyle,KeepTogether
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4,landscape
    from reportlab.lib import colors
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.graphics.shapes import Drawing,Rect,String
    from integrations.reporting.reportes import etiqueta,texto
    r=_datos(reporte); buffer=BytesIO()
    pagina=landscape(A4) if len(r['columnas'])>6 and r['titulo']!='Resumen general' else A4
    ancho=pagina[0]-64; nombre=obtener('general','nombre')
    doc=SimpleDocTemplate(buffer,pagesize=pagina,rightMargin=32,leftMargin=32,topMargin=74,bottomMargin=45,
        title='Reporte de '+r['titulo'],author=nombre,subject='Filtros: '+_filtros(r))
    def estilo(nombre,**kwargs):
        return ParagraphStyle(nombre,fontName='Helvetica',fontSize=9,leading=12,textColor=colors.HexColor(PALETA['texto']),**kwargs)
    normal=estilo('Cuerpo'); muted=estilo('Nota',spaceAfter=7); muted.textColor=colors.HexColor(PALETA['secundario'])
    titulo=estilo('Titulo',spaceAfter=10); titulo.fontSize=21; titulo.leading=25; titulo.fontName='Helvetica-Bold'
    seccion=estilo('Seccion',spaceBefore=12,spaceAfter=8); seccion.fontSize=12; seccion.fontName='Helvetica-Bold'
    celda=estilo('Celda'); numero=estilo('Numero',alignment=2)
    cabecera=estilo('Cabecera'); cabecera.fontName='Helvetica-Bold'; cabecera.textColor=colors.white
    def p(t,s=normal): return Paragraph(escape(str(t)).replace('\n','<br/>'),s)
    cuerpo=[p('Reporte de '+r['titulo'],titulo),p(obtener('reportes','encabezado'),muted),
        p('Generado: '+texto(r['generado']),muted),p('Filtros aplicados: '+_filtros(r),muted),Spacer(1,8)]
    # Resumen de cantidades y métricas ya calculadas desde las filas filtradas.
    valores=list(r['indicadores'].items())
    for start in range(0,len(valores),4):
        grupo=valores[start:start+4]
        cartas=[[p(_valor(k,v),seccion) for k,v in grupo],[p(k,muted) for k,v in grupo]]
        t=Table(cartas,colWidths=[ancho/len(grupo)]*len(grupo))
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor(PALETA['fondo'])),
            ('LINEABOVE',(0,0),(-1,0),2,colors.HexColor(PALETA['acento'])),
            ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),10),('BOTTOMPADDING',(0,1),(-1,1),8)]))
        cuerpo.extend([t,Spacer(1,10)])
    for nota in r.get('notas',[]): cuerpo.append(p(nota,muted))
    def tabla(columnas,filas):
        if not filas: return p('No hay información disponible para los filtros seleccionados.',muted)
        pesos=[ancho_columna(k) for k in columnas]
        for i,k in enumerate(columnas):
            if k in ('motivo','observacion') and not any(f.get(k) for f in filas): pesos[i]=1.1
            if k in ('ubicacion_origen','ubicacion_destino') and any(f.get(k) for f in filas): pesos[i]=2.0
        datos=[[p(etiqueta(k),cabecera) for k in columnas]]
        for fila in filas:
            datos.append([p(_valor(k,fila.get(k)),numero if k in NUMERICAS or isinstance(fila.get(k),(int,float,Decimal)) else celda) for k in columnas])
        t=LongTable(datos,colWidths=[ancho*v/sum(pesos) for v in pesos],repeatRows=1,hAlign='LEFT',splitInRow=1)
        reglas=[('BACKGROUND',(0,0),(-1,0),colors.HexColor(PALETA['primario'])),
            ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor(PALETA['fondo'])]),
            ('LINEBELOW',(0,0),(-1,0),1,colors.HexColor(PALETA['acento'])),
            ('INNERGRID',(0,0),(-1,-1),.3,colors.HexColor(PALETA['borde'])),
            ('BOX',(0,0),(-1,-1),.3,colors.HexColor(PALETA['borde'])),
            ('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),
            ('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)]
        for i,fila in enumerate(filas,1):
            for j,k in enumerate(columnas):
                if k in ('estado','estado_stock','nivel','riesgo'):
                    tonos=colores_estado(fila.get(k),k)
                    if tonos: reglas.append(('BACKGROUND',(j,i),(j,i),colors.HexColor('#'+tonos[0])))
        t.setStyle(TableStyle(reglas)); return t
    if r['titulo']!='Resumen general':
        cuerpo.append(p('Detalle · '+str(len(r['filas']))+' registros',seccion))
        cuerpo.append(tabla(r['columnas'],r['filas']))
    for s in r.get('secciones',[]):
        cuerpo.append(KeepTogether([p(s['titulo'],seccion),tabla(s['columnas'],s['filas'])]))
    if r['mensaje']: cuerpo.append(p(r['mensaje'],muted))
    totales=[(k,v) for k,v in r['totales'].items() if k!='Registros']
    if totales and r['titulo']!='Resumen general':
        encabezado=p('Totales por unidad' if any(k.startswith('Stock ') for k,v in totales) else 'Totales',seccion)
        t=Table([[p(k),p(_valor(k,v),numero)] for k,v in totales],colWidths=[ancho*.65,ancho*.35])
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor(PALETA['azul_claro'])),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7)])); cuerpo.append(KeepTogether([encabezado,t]))
    for grafico in r.get('graficos',[]):
        datos=grafico['datos']; alto=35+len(datos)*24
        dibujo=Drawing(ancho,alto); maximo=max(float(v) for _,v in datos) or 1
        izquierda=min(190,ancho*.37); largo=ancho-izquierda-95
        for i,(label,valor) in enumerate(datos):
            y=alto-25-i*24
            # Truncar sólo etiquetas del gráfico, manteniendo detalle completo en la tabla.
            label=str(label); label=label if len(label)<=32 else label[:29]+'...'
            dibujo.add(String(0,y,label,fontName='Helvetica',fontSize=9,fillColor=colors.HexColor(PALETA['texto'])))
            dibujo.add(Rect(izquierda,y-3,max(0,float(valor)/maximo*largo),13,fillColor=colors.HexColor(PALETA['acento']),strokeColor=None))
            dibujo.add(String(izquierda+largo+9,y,_valor('cantidad',valor),fontName='Helvetica',fontSize=9))
        cuerpo.append(KeepTogether([p(grafico['titulo']+' · '+grafico['unidad'],seccion),dibujo]))
    class Paginado(Canvas):
        def __init__(self,*a,**kw): super().__init__(*a,**kw); self.estados=[]
        def showPage(self): self.estados.append(dict(self.__dict__)); self._startPage()
        def save(self):
            total=len(self.estados)
            for estado in self.estados:
                self.__dict__.update(estado)
                self.setFillColor(colors.HexColor(PALETA['primario'])); self.rect(0,pagina[1]-49,pagina[0],49,fill=1,stroke=0)
                self.setFillColor(colors.white); self.setFont('Helvetica-Bold',10)
                marca=nombre if len(nombre)<=90 else nombre[:87]+'...'
                self.drawString(32,pagina[1]-28,marca)
                self.setStrokeColor(colors.HexColor(PALETA['borde'])); self.line(32,35,pagina[0]-32,35)
                self.setFont('Helvetica',8); self.setFillColor(colors.HexColor(PALETA['secundario']))
                self.drawString(32,22,marca[:48]+' | '+texto(r['generado']))
                self.drawRightString(pagina[0]-32,22,f'Página {self._pageNumber} de {total}')
                Canvas.showPage(self)
            Canvas.save(self)
    doc.build(cuerpo,canvasmaker=Paginado)
    return _guardar(buffer.getvalue(),archivo)

def exportar_excel(reporte,archivo):
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table,TableStyleInfo
    from openpyxl.worksheet.page import PageMargins
    from openpyxl.chart import BarChart,Reference
    from openpyxl.formatting.rule import FormulaRule
    from integrations.reporting.reportes import etiqueta,texto,TIPOS
    r=_datos(reporte); wb=Workbook(); wb.remove(wb.active)
    wb.properties.title='Reporte de '+r['titulo']; wb.properties.creator=obtener('general','nombre'); wb.properties.created=r['generado']
    wb.properties.subject='Filtros: '+_filtros(r)
    borde=Border(bottom=Side(style='hair',color=PALETA['borde'].strip('#')))
    def poner(ws,i,j,v,bold=False,fondo=None,color=None):
        if isinstance(v,Decimal): v=float(v)
        c=ws.cell(i,j,v)
        if isinstance(v,str): c.data_type='s'
        c.font=Font(name=FUENTE_EXCEL,size=11,bold=bold,color=(color or PALETA['texto']).strip('#'))
        c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='right' if isinstance(v,(int,float)) else 'left')
        c.border=borde
        if fondo: c.fill=PatternFill('solid',fgColor=fondo.strip('#'))
        if isinstance(v,datetime): c.number_format=FORMATOS['hora']
        elif isinstance(v,date): c.number_format=FORMATOS['fecha']
        elif isinstance(v,(int,float)): c.number_format=FORMATOS['entero'] if float(v).is_integer() else FORMATOS['cantidad']
        return c
    def fila_titulo(ws,i,text,n,principal=False):
        c=poner(ws,i,1,text,True,PALETA['primario'] if principal else PALETA['fondo'], '#FFFFFF' if principal else PALETA['primario'])
        ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=max(2,n))
        c.font=Font(name=FUENTE_EXCEL,size=18 if principal else 12,bold=True,color='FFFFFF' if principal else PALETA['primario'].strip('#'))
        ws.row_dimensions[i].height=max(30,math.ceil(len(text)/(max(2,n)*19))*18)
    def pagina(ws,n):
        ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape' if n>6 else 'portrait'
        ws.page_setup.paperSize=ws.PAPERSIZE_A4; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.page_margins=PageMargins(left=.3,right=.3,top=.45,bottom=.45,header=.2,footer=.2)
        ws.oddFooter.left.text=obtener('general','nombre').replace('&','&&')[:65]
        ws.oddFooter.center.text=r['generado'].strftime('%d/%m/%Y %H:%M')
        ws.oddFooter.right.text='Página &P de &N'
    def detalle(titulo,columnas,filas,totales=None):
        ws=wb.create_sheet(titulo); n=len(columnas); pagina(ws,n)
        fila_titulo(ws,1,obtener('general','nombre'),n,True)
        fila_titulo(ws,2,r['titulo'],n)
        fila_titulo(ws,3,'Generado: '+texto(r['generado']),n)
        fila_titulo(ws,4,'Filtros: '+_filtros(r),n)
        fila_titulo(ws,5,' · '.join(r.get('notas',[])) or obtener('reportes','encabezado'),n)
        for j,k in enumerate(columnas,1): poner(ws,6,j,etiqueta(k),True,PALETA['primario'],'#FFFFFF')
        ws.row_dimensions[6].height=32
        for j,k in enumerate(columnas,1):
            longitud=max([len(etiqueta(k))]+[min(70,len(str(f.get(k,'')))) for f in filas])
            ws.column_dimensions[get_column_letter(j)].width=max(12,min(48,longitud+3))
        for i,f in enumerate(filas,7):
            lineas=1
            for j,k in enumerate(columnas,1):
                v=f.get(k)
                if v is None: v=_valor(k,v)
                c=poner(ws,i,j,v,False,PALETA['fondo'] if i%2 else '#FFFFFF')
                if k in ('precio','valor'): c.number_format=FORMATOS['dinero']
                elif k in NUMERICAS and isinstance(v,(int,float,Decimal)): c.number_format=FORMATOS['cantidad']
                width=ws.column_dimensions[get_column_letter(j)].width
                lineas=max(lineas,sum(max(1,math.ceil(len(s)/(width*.88))) for s in str(v).split('\n')))
            ws.row_dimensions[i].height=min(409,max(25,lineas*16))
        fin=6+len(filas)
        if filas:
            t=Table(displayName='Datos'+str(len(wb.worksheets)),ref=f'A6:{get_column_letter(n)}{fin}')
            t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(t)
            for j,k in enumerate(columnas,1):
                if k not in ('estado_stock','riesgo','nivel','estado'): continue
                letra=get_column_letter(j)
                for estado in sorted({str(f.get(k,'')) for f in filas}):
                    tonos=colores_estado(estado,k)
                    if tonos:
                        regla=FormulaRule(formula=[f'${letra}7="{estado.replace(chr(34),chr(34)*2)}"'],fill=PatternFill('solid',fgColor=tonos[0]),font=Font(color=tonos[1],bold=True))
                        ws.conditional_formatting.add(f'{letra}7:{letra}{fin}',regla)
        else:
            fila_titulo(ws,7,r['mensaje'] or 'No hay información disponible para los filtros seleccionados.',n)
        pos=max(fin+3,10)
        for k,v in (r['totales'] if totales is None else totales).items():
            poner(ws,pos,1,k,True,PALETA['azul_claro']); poner(ws,pos,2,v,True,PALETA['azul_claro']); pos+=1
        ws.freeze_panes='A7'; ws.auto_filter.ref=f'A6:{get_column_letter(n)}{max(6,fin)}'
        ws.print_title_rows='1:6'; ws.print_options.horizontalCentered=True
        ws.print_area=f'A1:{get_column_letter(max(2,n))}{pos}'
        return ws
    complejo=r['titulo'] in TIPOS and (len(r['columnas'])>5 or r['titulo']=='Resumen general')
    if complejo:
        ws=wb.create_sheet('RESUMEN'); pagina(ws,10)
        for col in 'ABCDEFGHIJ': ws.column_dimensions[col].width=16
        fila_titulo(ws,1,obtener('general','nombre'),10,True)
        fila_titulo(ws,2,'Reporte de '+r['titulo'],10)
        fila_titulo(ws,3,'Generado: '+texto(r['generado']),10)
        fila_titulo(ws,4,'Filtros: '+_filtros(r),10)
        pos=6
        for k,v in r['indicadores'].items():
            poner(ws,pos,1,k,True,PALETA['fondo']); ws.merge_cells(start_row=pos,start_column=1,end_row=pos,end_column=4)
            c=poner(ws,pos,5,v,True,PALETA['azul_claro']); ws.merge_cells(start_row=pos,start_column=5,end_row=pos,end_column=6)
            if k=='Valor a precio registrado': c.number_format=FORMATOS['dinero']
            ws.row_dimensions[pos].height=28; pos+=1
        for k,v in r['totales'].items():
            if k.startswith('Stock '):
                poner(ws,pos,1,k,True); ws.merge_cells(start_row=pos,start_column=1,end_row=pos,end_column=4)
                poner(ws,pos,5,v); pos+=1
        pos+=2
        for nota in r.get('notas',[]):
            fila_titulo(ws,pos,nota,10); ws.row_dimensions[pos].height=max(35,math.ceil(len(nota)/130)*18); pos+=2
        for g in r.get('graficos',[]):
            fila_titulo(ws,pos,g['titulo']+' · '+g['unidad'],10)
            poner(ws,pos+2,1,'Grupo',True); poner(ws,pos+2,3,g['unidad'],True)
            for i,(k,v) in enumerate(g['datos'],pos+3):
                poner(ws,i,1,k); ws.merge_cells(start_row=i,start_column=1,end_row=i,end_column=2); poner(ws,i,3,v)
                ws.row_dimensions[i].height=30
            chart=BarChart(); chart.type='bar'; chart.title=g['titulo']; chart.y_axis.title=''; chart.x_axis.title=g['unidad']
            chart.add_data(Reference(ws,min_col=3,min_row=pos+2,max_row=pos+2+len(g['datos'])),titles_from_data=True)
            chart.set_categories(Reference(ws,min_col=1,min_row=pos+3,max_row=pos+2+len(g['datos'])))
            chart.legend=None; chart.width=19; chart.height=9
            chart.series[0].graphicalProperties.solidFill=PALETA['acento'].strip('#')
            chart.series[0].graphicalProperties.line.noFill=True
            ws.add_chart(chart,f'D{pos+2}'); pos+=max(23,len(g['datos'])+6)
        ws.freeze_panes='A6'; ws.print_area=f'A1:J{pos}'; ws.print_title_rows='1:4'
    if r['titulo']!='Resumen general': detalle('DETALLE' if complejo else 'Reporte',r['columnas'],r['filas'])
    for i,s in enumerate(r.get('secciones',[]),1): detalle(('Actividad','Críticos','Predicción','Anomalías','Estado')[i-1],s['columnas'],s['filas'],{'Registros':len(s['filas'])})
    buffer=BytesIO(); wb.save(buffer)
    return _guardar(buffer.getvalue(),archivo)
