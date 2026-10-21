"""Regresión: filas SQL temporales con rollback, exportadores y configuración aislada."""

import sys
from pathlib import Path as _RutaProyecto
_RAIZ_PROYECTO = _RutaProyecto(__file__).resolve().parents[1]
if str(_RAIZ_PROYECTO) not in sys.path:
    sys.path.insert(0, str(_RAIZ_PROYECTO))

import unittest
import tempfile
from pathlib import Path
from copy import deepcopy
from unittest.mock import patch,MagicMock
from datetime import datetime,timedelta
from decimal import Decimal
from probar_inventario import InventarioSQL,ConexionPrueba
from ai.anomalies import deteccion_anomalias as anom
from core import config_manager as cfg
from integrations.reporting import reportes

class OperacionesSQL(InventarioSQL):
    def setUp(self):
        super().setUp()
        p=patch.object(anom,'conectar',side_effect=lambda:ConexionPrueba(self.c,'anom_test')); p.start(); self.addCleanup(p.stop)
    def test_hallazgo_estados_y_reporte_comparten_datos(self):
        from inventory import movimientos
        self.entrada(1000)
        for _ in range(10): movimientos.registrar_salida(self.emp,self.p,self.a,5)
        m=movimientos.registrar_salida(self.emp,self.p,self.a,50)
        filas=anom.consultar_anomalias(producto=self.p)
        self.assertEqual(len(filas),1); a=filas[0]
        self.assertEqual(a['id_movimiento'],m); self.assertEqual(a['nivel'],'CRÍTICA')
        self.assertAlmostEqual(a['score'],10); self.assertIn('promedio previo 5.00',a['motivo'])
        anom.cambiar_estado_anomalia(a['id_anomalia'],'JUSTIFICADA','Pedido real confirmado')
        anom.consultar_anomalias(producto=self.p)
        filas=anom.listar_anomalias(producto=self.p)
        self.assertEqual(len(filas),1); self.assertEqual(filas[0]['estado'],'JUSTIFICADA')
        r=reportes.obtener_datos_reporte('Anomalías',{'producto':self.p,'nivel':'CRÍTICA'})
        self.assertEqual(r['filas'][0]['id_anomalia'],a['id_anomalia'])
        with self.assertRaises(ValueError): anom.cambiar_estado_anomalia(a['id_anomalia'],'JUSTIFICADA','')
    def test_filtrar_inventario_y_ubicacion(self):
        from inventory import movimientos
        self.entrada('10.5'); movimientos.registrar_transferencia(self.emp,self.p,self.a,self.b,'2.5')
        r=reportes.obtener_datos_reporte('Inventario actual',{'producto':self.p})
        self.assertEqual(len(r['filas']),1); self.assertEqual(r['filas'][0]['stock_actual'],Decimal('10.5'))
        r=reportes.obtener_datos_reporte('Inventario por ubicación',{'producto':self.p,'ubicacion':self.b})
        self.assertEqual(len(r['filas']),1); self.assertEqual(r['filas'][0]['cantidad'],Decimal('2.5'))
    def test_reportes_sin_historial_y_filtros_invalidos(self):
        r=reportes.obtener_datos_reporte('Reconocimiento facial')
        self.assertEqual(r['filas'],[]); self.assertIn('Sin historial',r['mensaje'])
        with self.assertRaises(ValueError): reportes.obtener_datos_reporte('Empleados',{'desde':'2026-10-01'})
        with self.assertRaises(ValueError): reportes.obtener_datos_reporte('Salidas',{'desde':'2026-10-03','hasta':'2026-10-01'})

class Estadisticas(unittest.TestCase):
    def datos(self,cantidades):
        return [{'id_detalle':i+1,'id_producto':1,'id_movimiento':i+1,'producto':'Fixture estadístico','tipo_movimiento':'Salida','cantidad':q,'fecha':datetime(2026,10,1)+timedelta(minutes=i)} for i,q in enumerate(cantidades)]
    def test_no_incluir_candidato_ni_mezclar_productos(self):
        filas=self.datos([10]*10+[1000]); r=anom.analizar_datos_movimientos(filas)
        self.assertEqual(r['anomalias_detectadas'],1); self.assertEqual(r['resultados'][0]['score'],100)
        filas[-1]['id_producto']=2
        self.assertEqual(anom.analizar_datos_movimientos(filas)['anomalias_detectadas'],0)
    def test_mismo_tipo_y_minimo_historial(self):
        filas=self.datos([10]*9+[1000]); r=anom.analizar_datos_movimientos(filas)
        self.assertEqual(r['movimientos_evaluados'],0); self.assertIn('insuficientes',r['mensaje'])
        filas=self.datos([10]*10+[1000]); filas[-1]['tipo_movimiento']='Transferencia'
        self.assertEqual(anom.analizar_datos_movimientos(filas)['anomalias_detectadas'],0)
    def test_sensibilidad_modifica_clasificacion(self):
        filas=self.datos([10]*10+[35]); c=deepcopy(cfg.DEFAULTS['anomalias'])
        self.assertEqual(anom.analizar_datos_movimientos(filas,c)['anomalias_detectadas'],1)
        c['sensibilidad']=4
        self.assertEqual(anom.analizar_datos_movimientos(filas,c)['anomalias_detectadas'],0)

class Configuracion(unittest.TestCase):
    def test_json_faltante_corrupto_y_escritura_atomica(self):
        with tempfile.TemporaryDirectory() as d,patch.object(cfg,'ARCHIVO',Path(d)/'config.json'):
            c=cfg.cargar_configuracion(); self.assertEqual(c['prediccion']['periodo'],30)
            c['jarvis']['volumen']=.5; cfg.guardar_configuracion(c)
            self.assertEqual(cfg.cargar_configuracion()['jarvis']['volumen'],.5)
            cfg.ARCHIVO.write_text('{roto',encoding='utf-8')
            self.assertEqual(cfg.cargar_configuracion()['jarvis']['volumen'],.8)
            cfg.ARCHIVO.write_text('[]',encoding='utf-8')
            self.assertEqual(cfg.cargar_configuracion()['jarvis']['volumen'],.8)
    def test_validacion_no_guardar_secretos(self):
        c=deepcopy(cfg.DEFAULTS); c['facial']['fuente_video']='rtsp://usuario:secreto@localhost/video'
        with self.assertRaises(ValueError): cfg.validar(c)
        c=deepcopy(cfg.DEFAULTS); c['jarvis']['volumen']=2
        with self.assertRaises(ValueError): cfg.validar(c)

class Exportaciones(unittest.TestCase):
    def test_excel_conserva_numeros_y_no_ejecuta_texto_sql(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            r={'titulo':'Fixture de exportación','generado':datetime(2026,10,3),'filtros':{},'columnas':('codigo','cantidad','fecha'),
               'filas':[{'codigo':'=1+1','cantidad':Decimal('2.5'),'fecha':datetime(2026,10,3)}],'totales':{'Registros':1},'mensaje':''}
            p=Path(d)/'prueba.xlsx'; reportes.exportar_excel(r,p); ws=load_workbook(p).active
            self.assertEqual(ws['A7'].data_type,'s'); self.assertEqual(ws['A7'].value,'=1+1')
            self.assertEqual(ws['B7'].value,2.5); self.assertEqual(ws['C7'].value,datetime(2026,10,3))
            self.assertEqual(ws.freeze_panes,'A7')
    def test_pdf_paginacion_y_encabezados_repetidos(self):
        from pypdf import PdfReader
        with tempfile.TemporaryDirectory() as d:
            r={'titulo':'Fixture de paginación','generado':datetime(2026,10,3),'filtros':{},'columnas':('codigo','motivo'),
               'filas':[{'codigo':'Prueba '+str(i),'motivo':'Texto de prueba largo para verificar la paginación. '*4} for i in range(80)],'totales':{'Registros':80},'mensaje':''}
            p=Path(d)/'prueba.pdf'; reportes.exportar_pdf(r,p); pdf=PdfReader(p)
            self.assertGreater(len(pdf.pages),1)
            for page in pdf.pages: self.assertIn('Código',page.extract_text()); self.assertIn('Página',page.extract_text())

class Diagnosticos(unittest.TestCase):
    def test_fallo_sql_no_declara_online_ni_expone_detalle_sensible(self):
        from core import estado_sistema as estado
        with patch.dict(estado.ACCIONES,{'Base de datos':MagicMock(side_effect=RuntimeError('password=no_mostrar'))}),patch('traceback.print_exc'):
            r=estado.comprobar('Base de datos')
        self.assertEqual(r['estado'],'ERROR'); self.assertNotIn('password',r['detalle'])
        self.assertIn('comprobado',r)
        estado._pruebas.pop('Base de datos',None)
    def test_camara_sin_imagen_libera_recurso_y_no_declara_disponible(self):
        from core import estado_sistema as estado
        from ai.facial.reconocimiento_cnn import _bloqueo_camara
        cam=MagicMock(); cam.read.return_value=(False,None)
        with patch('surveillance.camara_utils.abrir_camara',return_value=cam):
            with self.assertRaises(RuntimeError): estado.probar_camara()
        cam.release.assert_called_once(); self.assertFalse(_bloqueo_camara.locked())

def load_tests(loader,tests,pattern):
    suite=unittest.TestSuite()
    for nombre in ('test_hallazgo_estados_y_reporte_comparten_datos','test_filtrar_inventario_y_ubicacion','test_reportes_sin_historial_y_filtros_invalidos'): suite.addTest(OperacionesSQL(nombre))
    for clase in (Estadisticas,Configuracion,Exportaciones,Diagnosticos): suite.addTests(loader.loadTestsFromTestCase(clase))
    return suite
