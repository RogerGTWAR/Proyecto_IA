"""Pruebas de instantáneas, unidades, tipos Excel y exportaciones sin sobrescritura."""

import sys
from pathlib import Path as _RutaProyecto
_RAIZ_PROYECTO = _RutaProyecto(__file__).resolve().parents[1]
if str(_RAIZ_PROYECTO) not in sys.path:
    sys.path.insert(0, str(_RAIZ_PROYECTO))

import unittest
from tempfile import TemporaryDirectory
from pathlib import Path
from datetime import datetime
from decimal import Decimal
from unittest.mock import patch
from openpyxl import load_workbook
from pypdf import PdfReader
from integrations.reporting import reportes

class Reportes(unittest.TestCase):
    def reporte(self):
        return reportes.preparar_reporte({'titulo':'Inventario actual','generado':datetime(2026,10,4,10,30),
            'filtros':{},'columnas':('codigo','nombre','categoria','unidad_medida','stock_actual','stock_minimo','estado_stock','ubicaciones'),
            'filas':[{'codigo':'=SQL','nombre':'Filtro','categoria':'Repuestos','unidad_medida':'Unidad','stock_actual':Decimal('20.5'),'stock_minimo':10,'estado_stock':'NORMAL','ubicaciones':'A; B','precio':Decimal('25')},
                     {'codigo':'ACE','nombre':'Aceite','categoria':'Lubricantes','unidad_medida':'Litro','stock_actual':Decimal('2.125'),'stock_minimo':5,'estado_stock':'CRÍTICO','ubicaciones':'C','precio':Decimal('100')}],
            'totales':{'Stock Unidad':Decimal('20.5'),'Stock Litro':Decimal('2.125'),'Registros':2},'mensaje':''})
    def test_snapshot_exports_do_not_query_or_modify(self):
        r=self.reporte(); from copy import deepcopy
        antes=deepcopy(r)
        with TemporaryDirectory() as d,patch.object(reportes,'consultar',side_effect=AssertionError('SQL durante exportación')):
            reportes.exportar_pdf(r,Path(d)/'inventario.pdf'); reportes.exportar_excel(r,Path(d)/'inventario.xlsx')
        self.assertEqual(r,antes)
    def test_excel_typed_tables_print_charts_and_conditional_states(self):
        r=self.reporte()
        with TemporaryDirectory() as d:
            archivo=reportes.exportar_excel(r,Path(d)/'inventario.xlsx'); wb=load_workbook(archivo)
            self.assertEqual(wb.sheetnames,['RESUMEN','DETALLE'])
            ws=wb['DETALLE']; self.assertEqual(ws['A7'].value,'=SQL'); self.assertEqual(ws['A7'].data_type,'s')
            self.assertEqual(ws['E8'].value,2.125); self.assertEqual(ws['E8'].number_format,'#,##0.00')
            self.assertEqual(ws.freeze_panes,'A7'); self.assertTrue(ws.tables); self.assertTrue(ws.conditional_formatting)
            self.assertEqual(ws.page_setup.fitToWidth,1); self.assertTrue(ws.print_title_rows)
            self.assertEqual(len(wb['RESUMEN']._charts),2)
            self.assertEqual(wb.properties.created,datetime(2026,10,4,10,30))
    def test_units_and_stock_not_duplicated(self):
        r=self.reporte()
        self.assertEqual(r['totales']['Stock Unidad'],Decimal('20.5'))
        self.assertEqual(r['totales']['Stock Litro'],Decimal('2.125'))
        self.assertEqual(r['indicadores']['Valor a precio registrado'],Decimal('725.0'))
        self.assertEqual({g['unidad'] for g in r['graficos']},{'Unidad','Litro'})
    def test_no_silent_overwrite(self):
        with TemporaryDirectory() as d:
            p=Path(d)/'reporte.pdf'; p.write_bytes(b'original')
            q=reportes.exportar_pdf(self.reporte(),p)
            self.assertEqual(q.name,'reporte_01.pdf'); self.assertEqual(p.read_bytes(),b'original')
    def test_empty_and_pagination(self):
        r={'titulo':'Reconocimiento facial','generado':datetime(2026,10,4),'filtros':{},'columnas':('fecha','persona','resultado'),
           'filas':[],'totales':{'Registros':0},'mensaje':'Sin historial persistente de reconocimiento facial.'}
        with TemporaryDirectory() as d:
            p=reportes.exportar_pdf(r,Path(d)/'vacio.pdf')
            t=''.join(x.extract_text() for x in PdfReader(p).pages)
            self.assertIn('Sin historial persistente',t)
            r=self.reporte(); r['filas']=r['filas']*60; r=reportes.preparar_reporte(r)
            p=reportes.exportar_pdf(r,Path(d)/'largo.pdf'); pdf=PdfReader(p)
            self.assertGreater(len(pdf.pages),1)
            for i,page in enumerate(pdf.pages,1):
                self.assertIn(f'Página {i} de {len(pdf.pages)}',page.extract_text())
    def test_date_filters_rejected(self):
        with self.assertRaises(ValueError): reportes.obtener_datos_reporte('Salidas',{'desde':'2026-10-04','hasta':'2026-10-01'})
    def test_state_risk_low_has_different_meaning(self):
        from integrations.reporting.report_theme import colores_estado
        self.assertNotEqual(colores_estado('BAJO','estado_stock'),colores_estado('BAJO','riesgo'))

if __name__=='__main__': unittest.main()
